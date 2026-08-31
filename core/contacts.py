"""Leitura de listas de contatos (CSV, XLSX ou colado) e validação de telefones."""

import csv
import io
import re
from typing import Dict, List, Optional, Tuple

# DDDs válidos no Brasil (Anatel).
DDDS_VALIDOS = {
    11, 12, 13, 14, 15, 16, 17, 18, 19, 21, 22, 24, 27, 28, 31, 32, 33, 34, 35, 37, 38,
    41, 42, 43, 44, 45, 46, 47, 48, 49, 51, 53, 54, 55, 61, 62, 63, 64, 65, 66, 67, 68,
    69, 71, 73, 74, 75, 77, 79, 81, 82, 83, 84, 85, 86, 87, 88, 89, 91, 92, 93, 94, 95,
    96, 97, 98, 99,
}

PISTAS_TELEFONE = ("telefone", "phone", "celular", "whatsapp", "fone", "tel", "numero", "número", "contato")


def _somente_digitos(valor: str) -> str:
    return "".join(c for c in str(valor) if c.isdigit())


def normalize_phone(bruto: str, adicionar_nono_digito: bool = False) -> Dict:
    """
    Normaliza um telefone brasileiro para o formato E.164 sem '+'.

    Devolve dict com: original, phone, status (ok|warn|invalid), reason.
    """
    original = str(bruto or "").strip()
    internacional = original.startswith("+") and not original.startswith("+55")
    digitos = _somente_digitos(original)

    if not digitos:
        return {"original": original, "phone": "", "status": "invalid", "reason": "Vazio"}

    if internacional:
        if 8 <= len(digitos) <= 15:
            return {"original": original, "phone": digitos, "status": "warn", "reason": "Número internacional (não validado)"}
        return {"original": original, "phone": "", "status": "invalid", "reason": "Número internacional fora do padrão E.164"}

    # Descobre se o código do país já veio junto, olhando pelo comprimento total.
    if len(digitos) in (12, 13) and digitos.startswith("55"):
        nacional = digitos[2:]
    elif len(digitos) in (10, 11):
        nacional = digitos
    else:
        return {
            "original": original,
            "phone": "",
            "status": "invalid",
            "reason": f"{len(digitos)} dígitos — esperado 10 ou 11 (com DDD)",
        }

    ddd = int(nacional[:2])
    if ddd not in DDDS_VALIDOS:
        return {"original": original, "phone": "", "status": "invalid", "reason": f"DDD {ddd:02d} não existe"}

    assinante = nacional[2:]
    if len(set(assinante)) == 1:
        return {"original": original, "phone": "", "status": "invalid", "reason": "Número repetitivo (ex: 000000000)"}

    reason = ""
    status = "ok"

    if len(assinante) == 8:
        if assinante[0] in "6789":
            # Celular antigo, gravado sem o nono dígito.
            if adicionar_nono_digito:
                assinante = "9" + assinante
                status = "ok"
                reason = "9º dígito adicionado automaticamente"
            else:
                status = "warn"
                reason = "Celular sem o 9º dígito — pode não ser entregue"
        else:
            status = "warn"
            reason = "Parece telefone fixo — WhatsApp provavelmente não entrega"
    elif len(assinante) == 9 and assinante[0] != "9":
        status = "warn"
        reason = "Celular de 9 dígitos que não começa com 9"

    return {"original": original, "phone": f"55{ddd:02d}{assinante}", "status": status, "reason": reason}


# ---------------------------------------------------------------- leitura de arquivos


def _decodificar(conteudo: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return conteudo.decode(encoding)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("utf-8", errors="replace")


def parse_csv(conteudo: bytes) -> Tuple[List[str], List[Dict]]:
    texto = _decodificar(conteudo)
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t|")
        delimitador = dialeto.delimiter
    except csv.Error:
        delimitador = ";" if amostra.count(";") > amostra.count(",") else ","

    leitor = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    colunas = [c.strip() for c in (leitor.fieldnames or []) if c and c.strip()]
    linhas = []
    for registro in leitor:
        linhas.append({(k or "").strip(): (v or "").strip() for k, v in registro.items() if k})
    return colunas, linhas


def parse_xlsx(conteudo: bytes) -> Tuple[List[str], List[Dict]]:
    from openpyxl import load_workbook

    planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    aba = planilha.active
    iterador = aba.iter_rows(values_only=True)

    try:
        cabecalho = next(iterador)
    except StopIteration:
        return [], []

    colunas = [str(c).strip() if c is not None else f"coluna_{i+1}" for i, c in enumerate(cabecalho)]
    linhas = []
    for valores in iterador:
        if all(v is None or str(v).strip() == "" for v in valores):
            continue
        registro = {}
        for indice, coluna in enumerate(colunas):
            valor = valores[indice] if indice < len(valores) else None
            if isinstance(valor, float) and valor.is_integer():
                valor = int(valor)  # evita telefones virarem 5.4999674344e+10
            registro[coluna] = "" if valor is None else str(valor).strip()
        linhas.append(registro)
    planilha.close()
    return colunas, linhas


def parse_pasted(texto: str) -> Tuple[List[str], List[Dict]]:
    """Interpreta uma lista colada — um telefone por linha (separadores comuns aceitos)."""
    linhas = []
    for parte in re.split(r"[\n\r;,]+", texto or ""):
        parte = parte.strip()
        if parte:
            linhas.append({"telefone": parte})
    return ["telefone"], linhas


def parse_upload(nome_arquivo: str, conteudo: bytes) -> Tuple[List[str], List[Dict]]:
    nome = (nome_arquivo or "").lower()
    if nome.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(conteudo)
    if nome.endswith(".xls"):
        raise ValueError("Formato .xls antigo não é suportado — salve como .xlsx ou .csv")
    return parse_csv(conteudo)


def guess_phone_column(colunas: List[str]) -> Optional[str]:
    """Escolhe a coluna que mais parece conter telefone."""
    for coluna in colunas:
        if any(pista in coluna.lower() for pista in PISTAS_TELEFONE):
            return coluna
    return colunas[0] if colunas else None


# ---------------------------------------------------------------- montagem da lista


def build_recipients(
    linhas: List[Dict],
    coluna_telefone: str,
    mapeamento: Dict[str, Dict],
    adicionar_nono_digito: bool = False,
    blocklist: Optional[set] = None,
) -> Dict:
    """
    Constrói a lista final de destinatários já validada.

    `mapeamento` liga a chave da variável do template a uma origem:
    {"body:1": {"type": "column", "value": "nome"}} ou {"type": "fixed", "value": "texto"}
    """
    blocklist = blocklist or set()
    validos, avisos, invalidos, duplicados, bloqueados = [], [], [], [], []
    vistos = set()

    for indice, linha in enumerate(linhas, start=1):
        bruto = linha.get(coluna_telefone, "")
        resultado = normalize_phone(bruto, adicionar_nono_digito)
        registro = {
            "row": indice,
            "original": resultado["original"],
            "phone": resultado["phone"],
            "status": resultado["status"],
            "reason": resultado["reason"],
            "values": {},
        }

        for chave, origem in (mapeamento or {}).items():
            if origem.get("type") == "column":
                registro["values"][chave] = str(linha.get(origem.get("value", ""), "") or "")
            else:
                registro["values"][chave] = str(origem.get("value", "") or "")

        if resultado["status"] == "invalid":
            invalidos.append(registro)
            continue
        if resultado["phone"] in blocklist:
            registro["reason"] = "Na blocklist (opt-out)"
            bloqueados.append(registro)
            continue
        if resultado["phone"] in vistos:
            registro["reason"] = "Telefone duplicado na lista"
            duplicados.append(registro)
            continue

        vistos.add(resultado["phone"])
        if resultado["status"] == "warn":
            avisos.append(registro)
        validos.append(registro)

    return {
        "recipients": validos,
        "summary": {
            "total_linhas": len(linhas),
            "validos": len(validos),
            "com_aviso": len(avisos),
            "invalidos": len(invalidos),
            "duplicados": len(duplicados),
            "bloqueados": len(bloqueados),
        },
        "warnings": avisos[:50],
        "invalid": invalidos[:50],
        "duplicates": duplicados[:50],
        "blocked": bloqueados[:50],
    }
